#!/usr/bin/env python3
"""
docx-export.py — Export Plain English markdown to Word (.docx)

Usage from terminal:
  python docx-export.py input.md output.docx

Usage from execute_code:
  import subprocess
  subprocess.run(["python", "docx-export.py", "contract.md", "contract.docx"])

Requires: pip install python-docx
"""

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
import re
import sys

def md_to_docx(md_path, docx_path):
    with open(md_path, "r", encoding="utf-8") as f:
        content = f.read()

    doc = Document()
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(11)
    style.element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')

    lines = content.split('\n')
    i = 0

    while i < len(lines):
        line = lines[i].rstrip()
        if not line:
            i += 1
            continue

        if line.startswith('---'):
            i += 1
            continue

        if line.startswith('# ') and not line.startswith('## '):
            doc.add_heading(line[2:].strip(), level=1)
            i += 1
            continue

        if line.startswith('## '):
            doc.add_heading(line[3:].strip(), level=2)
            i += 1
            continue

        if line.startswith('### '):
            doc.add_heading(line[4:].strip(), level=3)
            i += 1
            continue

        # Table
        if line.startswith('|') and i + 1 < len(lines) and '|-' in lines[i+1]:
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                table_lines.append(lines[i].strip())
                i += 1

            rows = []
            for tl in table_lines:
                if re.match(r'^\|[\s\-|]+\|$', tl):
                    continue
                cells = [cell.strip() for cell in tl.split('|')[1:-1]]
                if cells:
                    rows.append(cells)

            if rows:
                num_cols = max(len(r) for r in rows)
                table = doc.add_table(rows=len(rows), cols=num_cols)
                table.style = 'Table Grid'
                for row_idx, row_cells in enumerate(rows):
                    for col_idx, cell_text in enumerate(row_cells):
                        if col_idx < num_cols:
                            cell = table.rows[row_idx].cells[col_idx]
                            cell.text = cell_text
                            if row_idx == 0:
                                for paragraph in cell.paragraphs:
                                    for run in paragraph.runs:
                                        run.bold = True
            continue

        # Bullet list
        if line.startswith('- '):
            p = doc.add_paragraph(style='List Bullet')
            text = line[2:].strip()
            parts = re.split(r'(\*\*.*?\*\*)', text)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    run = p.add_run(part[2:-2])
                    run.bold = True
                else:
                    p.add_run(part)
            i += 1
            continue

        # Regular paragraph with bold parsing
        p = doc.add_paragraph()
        parts = re.split(r'(\*\*.*?\*\*)', line)
        for part in parts:
            if part.startswith('**') and part.endswith('**'):
                run = p.add_run(part[2:-2])
                run.bold = True
            else:
                p.add_run(part)
        i += 1

    doc.save(docx_path)
    print(f"Saved: {docx_path}")

if __name__ == "__main__":
    if len(sys.argv) != 3:
        print("Usage: python docx-export.py input.md output.docx")
        sys.exit(1)
    md_to_docx(sys.argv[1], sys.argv[2])
